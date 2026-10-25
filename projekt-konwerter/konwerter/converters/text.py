"""Grupa B – tekst i dokumenty: TXT -> DOCX i DOCX -> TXT (biblioteka python-docx).

Ustalona semantyka (opisana też w README):

TXT -> DOCX
- plik TXT w UTF-8 (z BOM lub bez),
- KAŻDA linia tekstu = jeden akapit w DOCX, pusta linia = pusty akapit
  (dzięki temu układ tekstu zostaje taki sam, a konwersja w obie strony jest bezstratna),
- tabulatory zostają tabulatorami; znaki sterujące niedozwolone w XML (np. \\x00) są usuwane,
- czcionka i style: domyślne z szablonu python-docx (Calibri 11).

DOCX -> TXT
- tekst akapitów i tabel w kolejności z dokumentu,
- komórki tabeli rozdzielone tabulatorem, każdy wiersz tabeli w osobnej linii,
- formatowanie (pogrubienia, kolory, obrazy, przypisy, nagłówki/stopki strony) jest pomijane,
- wynik w UTF-8 bez BOM.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

from ..core import converter
from ..errors import InvalidInputError, MissingDependencyError
from .common import read_text_utf8

# Znaki sterujące, których nie wolno zapisać w XML (dozwolone są tylko \t, \n, \r).
_XML_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _import_docx():
    """Importuje python-docx dopiero przy konwersji – brak biblioteki to kontrolowany błąd,
    a nie awaria całej aplikacji przy starcie."""
    try:
        import docx
    except ImportError as exc:
        raise MissingDependencyError(
            "Brak biblioteki python-docx potrzebnej do plików DOCX.",
            "Zainstaluj zależności: python -m pip install -r requirements.txt",
        ) from exc
    return docx


@converter("txt", "docx")
def txt_to_docx(source: Path, target: Path) -> None:
    docx = _import_docx()
    text = read_text_utf8(source)
    text = _XML_ILLEGAL.sub("", text)

    document = docx.Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    document.save(str(target))


def _open_docx(source: Path):
    docx = _import_docx()
    from docx.opc.exceptions import PackageNotFoundError

    try:
        return docx.Document(str(source))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise InvalidInputError(
            f"Plik {source.name} nie jest poprawnym dokumentem DOCX albo jest uszkodzony.",
            "Otwórz go w Wordzie/LibreOffice i zapisz ponownie jako .docx (stary format .doc nie jest obsługiwany).",
        ) from exc


@converter("docx", "txt")
def docx_to_txt(source: Path, target: Path) -> None:
    from docx.table import Table

    document = _open_docx(source)

    lines: list[str] = []
    # iter_inner_content() zwraca akapity i tabele w kolejności, w jakiej są w dokumencie
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                lines.append("\t".join(cell.text for cell in row.cells))
        else:
            lines.append(block.text)

    with target.open("w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lines))
        f.write("\n")
